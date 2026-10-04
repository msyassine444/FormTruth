from typing import Optional
from io import BytesIO

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from .base import BaseExtractor, ExtractionError


class DocxExtractor(BaseExtractor):
    SUPPORTED_EXTENSIONS = {".docx"}
    SUPPORTED_CONTENT_TYPES = {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/octet-stream",
        "application/msword",  # بعض المتصفحات ترسله خطأً
    }

    @staticmethod
    def _paragraphs_text(doc) -> list[str]:
        return [p.text for p in doc.paragraphs if p.text and p.text.strip()]

    @staticmethod
    def _tables_text(doc) -> list[str]:
        rows: list[str] = []
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                if cells:
                    rows.append(" | ".join(cells))
        return rows

    def extract(
        self,
        filename: Optional[str],
        content_type: Optional[str],
        data: bytes,
    ) -> str:
        self.validate(filename, content_type, data)

        try:
            doc = Document(BytesIO(data))
        except PackageNotFoundError as e:
            raise ExtractionError(f"Invalid DOCX: {e}") from e
        except Exception as e:
            raise ExtractionError(f"Failed to open DOCX: {e}") from e

        parts: list[str] = []
        parts.extend(self._paragraphs_text(doc))
        parts.extend(self._tables_text(doc))

        full_text = "\n".join(parts).strip()
        if not full_text:
            raise ExtractionError("No extractable text found in DOCX")
        return full_text


_extractor = DocxExtractor()


def extract_text(
    filename: Optional[str],
    content_type: Optional[str],
    data: bytes,
) -> str:
    return _extractor.extract(filename, content_type, data)
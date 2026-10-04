from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient

from backend.main import app
from backend.extractors.docx_extractor import extract_text

client = TestClient(app)


def _make_docx(text: str) -> bytes:
    doc = Document()
    for line in text.split("\n"):
        doc.add_paragraph(line)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_docx_with_table(rows: list[list[str]]) -> bytes:
    doc = Document()
    doc.add_paragraph("CV")
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            table.cell(i, j).text = cell
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------- Extractor direct ----------

def test_docx_extracts_paragraphs():
    data = _make_docx("Name: Yassine Annous\nEmail: yassine@example.com")
    text = extract_text("cv.docx", "application/octet-stream", data)
    assert "Yassine Annous" in text
    assert "yassine@example.com" in text


def test_docx_extracts_table_cells():
    data = _make_docx_with_table([
        ["Field", "Value"],
        ["Name", "Yassine Annous"],
        ["Email", "yassine@example.com"],
    ])
    text = extract_text("cv.docx", "application/octet-stream", data)
    assert "Yassine Annous" in text
    assert "yassine@example.com" in text


def test_docx_rejects_invalid_bytes():
    with pytest.raises(Exception):
        extract_text("file.docx", "application/octet-stream", b"not a docx")


def test_docx_rejects_empty():
    with pytest.raises(Exception):
        extract_text("file.docx", "application/octet-stream", b"")


def test_docx_rejects_empty_document():
    doc = Document()
    buf = BytesIO()
    doc.save(buf)
    with pytest.raises(Exception):
        extract_text("empty.docx", "application/octet-stream", buf.getvalue())


# ---------- Endpoint ----------

def test_extract_endpoint_docx():
    data = _make_docx("Name: Yassine Annous\nDOB: 2009-05-12")
    response = client.post(
        "/extract",
        files={"file": ("cv.docx", data, "application/octet-stream")},
    )
    assert response.status_code == 200
    body = response.json()
    assert body["filename"] == "cv.docx"
    assert "Yassine Annous" in body["text"]


def test_extract_facts_endpoint_docx():
    data = _make_docx(
        "Name: Yassine Annous\n"
        "Email: yassine@example.com\n"
        "Phone: 0612345678\n"
        "DOB: 2009-05-12"
    )
    response = client.post(
        "/extract-facts",
        files={"file": ("cv.docx", data, "application/octet-stream")},
    )
    assert response.status_code == 200
    facts = response.json()["facts"]
    assert "yassine@example.com" in facts["emails"]
    assert "612345678" in facts["phones"]
    assert "2009-05-12" in facts["dates"]
    assert "Yassine Annous" in facts["names"]


# ---------- Full pipeline ----------

def test_full_pipeline_with_docx():
    data = _make_docx(
        "Name: Yassine Annous\n"
        "Email: yassine@example.com\n"
        "Phone: 0612345678\n"
        "DOB: 2009-05-12"
    )

    r1 = client.post(
        "/extract-facts",
        files={"file": ("cv.docx", data, "application/octet-stream")},
    )
    assert r1.status_code == 200
    facts = r1.json()["facts"]

    r2 = client.post("/build-truth", json={"facts": facts, "hints": {}})
    assert r2.status_code == 200
    truth = r2.json()["truth"]

    r3 = client.post("/compare", json={
        "truth": truth,
        "form": {
            "name": "Yassine Annous",
            "email": "yassine@example.com",
            "phone": "+212612345678",
            "dob": "12/05/2009",
        },
    })
    assert r3.status_code == 200
    assert r3.json()["summary"] == {"MATCH": 4, "CONFLICT": 0, "UNKNOWN": 0}